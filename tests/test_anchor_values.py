from __future__ import annotations

from datetime import date
from pathlib import Path
from typing import Any

import pytest
from openpyxl import Workbook

import excel_data_reader
from excel_data_reader import (
    DiagnosticCode,
    ExcelDataReaderError,
    ExcelReader,
    FormulaValue,
    ValueMode,
)


def _save_anchor_workbook(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Form"
    sheet["D7"] = "My_Cell"
    sheet["F6"] = 42
    workbook.save(path)
    workbook.close()


def test_get_value_from_anchor_applies_signed_row_and_column_offsets(tmp_path: Path) -> None:
    path = tmp_path / "anchors.xlsx"
    _save_anchor_workbook(path)

    with ExcelReader.open(path) as reader:
        match = reader.get_value_from_anchor(
            "my-cell",
            row_offset=-1,
            column_offset=2,
            sheet="Form",
            within="A1:H20",
        )

    assert match.anchor_cell.address == "D7"
    assert match.value_cell.address == "F6"
    assert match.value == 42
    assert isinstance(match, excel_data_reader.AnchorValueMatch)


def test_find_values_from_anchor_preserves_repeated_anchors_as_ambiguity(tmp_path: Path) -> None:
    path = tmp_path / "repeated-anchors.xlsx"
    workbook = Workbook()
    first = workbook.active
    first.title = "First"
    first["B2"] = "Account Total"
    first["C2"] = 10
    first["B5"] = "account_total"
    first["C5"] = 20
    second = workbook.create_sheet("Second")
    second["D3"] = "ACCOUNT-TOTAL"
    second["E3"] = 30
    workbook.save(path)
    workbook.close()

    with ExcelReader.open(path) as reader:
        matches = reader.find_values_from_anchor("account total")

    assert [match.anchor_cell.address for match in matches.matches] == ["B2", "B5", "D3"]
    assert [match.value for match in matches.matches] == [10, 20, 30]
    with pytest.raises(ExcelDataReaderError) as captured:
        matches.require_one()
    assert captured.value.diagnostics[0].code is DiagnosticCode.AMBIGUOUS_ANCHOR


def test_within_requires_the_offset_target_to_remain_inside_the_range(tmp_path: Path) -> None:
    path = tmp_path / "bounded-anchor.xlsx"
    _save_anchor_workbook(path)

    with ExcelReader.open(path) as reader:
        matches = reader.find_values_from_anchor(
            "my cell",
            row_offset=-1,
            column_offset=2,
            sheet="Form",
            within="A1:E20",
        )

    assert not matches.matches
    assert matches.diagnostics[0].code is DiagnosticCode.ANCHOR_TARGET_OUT_OF_BOUNDS
    assert matches.diagnostics[0].sheet == "Form"
    assert matches.diagnostics[0].address == "D7"


@pytest.mark.parametrize(
    ("anchor_value", "row_offset", "column_offset"),
    [
        ("", 0, 1),
        ("mycell", True, 1),
        ("mycell", 0, 1.5),
        ("mycell", 0, 0),
    ],
)
def test_anchor_lookup_rejects_invalid_anchor_and_offset_arguments(
    tmp_path: Path,
    anchor_value: str,
    row_offset: Any,
    column_offset: Any,
) -> None:
    path = tmp_path / "invalid-anchor-query.xlsx"
    _save_anchor_workbook(path)

    with (
        ExcelReader.open(path) as reader,
        pytest.raises(ExcelDataReaderError) as captured,
    ):
        reader.find_values_from_anchor(
            anchor_value,
            row_offset=row_offset,
            column_offset=column_offset,
        )

    assert captured.value.diagnostics[0].code is DiagnosticCode.INVALID_ANCHOR_QUERY


@pytest.mark.parametrize(
    ("row_offset", "column_offset"),
    [(-7, 0), (0, 16_381)],
)
def test_anchor_lookup_reports_targets_outside_the_excel_grid(
    tmp_path: Path,
    row_offset: int,
    column_offset: int,
) -> None:
    path = tmp_path / "grid-boundary.xlsx"
    _save_anchor_workbook(path)

    with ExcelReader.open(path) as reader:
        matches = reader.find_values_from_anchor(
            "my cell",
            row_offset=row_offset,
            column_offset=column_offset,
            sheet="Form",
        )

    assert not matches.matches
    assert matches.diagnostics[0].code is DiagnosticCode.ANCHOR_TARGET_OUT_OF_BOUNDS
    assert matches.diagnostics[0].address == "D7"


def test_workbook_wide_lookup_keeps_other_sheets_when_one_exceeds_scan_limit(
    tmp_path: Path,
) -> None:
    path = tmp_path / "bounded-workbook-scan.xlsx"
    workbook = Workbook()
    oversized = workbook.active
    oversized.title = "Oversized"
    oversized["Z100"] = "far away"
    compact = workbook.create_sheet("Compact")
    compact["A1"] = "Lookup"
    compact["B1"] = 99
    workbook.save(path)
    workbook.close()

    with ExcelReader.open(path, max_scan_cells=10) as reader:
        matches = reader.find_values_from_anchor("lookup")

    assert [match.value for match in matches.matches] == [99]
    assert matches.matches[0].anchor_cell.sheet == "Compact"
    assert matches.diagnostics[0].code is DiagnosticCode.SCAN_LIMIT_EXCEEDED
    assert matches.diagnostics[0].sheet == "Oversized"
    with pytest.raises(ExcelDataReaderError) as captured:
        matches.require_one()
    assert captured.value.diagnostics[0].code is DiagnosticCode.SCAN_LIMIT_EXCEEDED


def test_anchor_lookup_honors_the_reader_candidate_limit(tmp_path: Path) -> None:
    path = tmp_path / "many-anchors.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    for row, value in enumerate((10, 20, 30), start=1):
        sheet.cell(row, 1, "Lookup")
        sheet.cell(row, 2, value)
    workbook.save(path)
    workbook.close()

    with ExcelReader.open(path, max_candidates=2) as reader:
        matches = reader.find_values_from_anchor("lookup")

    assert [match.value for match in matches.matches] == [10, 20]
    assert matches.diagnostics[0].code is DiagnosticCode.SCAN_LIMIT_EXCEEDED


def test_candidate_limit_also_bounds_out_of_range_anchor_diagnostics(tmp_path: Path) -> None:
    path = tmp_path / "many-invalid-targets.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    for row in range(1, 4):
        sheet.cell(row, 1, "Lookup")
    workbook.save(path)
    workbook.close()

    with ExcelReader.open(path, max_candidates=2) as reader:
        matches = reader.find_values_from_anchor("lookup", within="A1:A3")

    assert not matches.matches
    assert [item.code for item in matches.diagnostics] == [
        DiagnosticCode.ANCHOR_TARGET_OUT_OF_BOUNDS,
        DiagnosticCode.ANCHOR_TARGET_OUT_OF_BOUNDS,
        DiagnosticCode.SCAN_LIMIT_EXCEEDED,
    ]


def test_default_offset_returns_the_blank_cell_immediately_right_of_anchor(
    tmp_path: Path,
) -> None:
    path = tmp_path / "blank-target.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "Lookup"
    workbook.save(path)
    workbook.close()

    with ExcelReader.open(path) as reader:
        match = reader.get_value_from_anchor("lookup")

    assert match.anchor_cell.address == "A1"
    assert match.value_cell.address == "B1"
    assert match.value is None


def test_sheet_and_within_disambiguate_repeated_anchors(tmp_path: Path) -> None:
    path = tmp_path / "scoped-anchors.xlsx"
    workbook = Workbook()
    first = workbook.active
    first.title = "First"
    first["B2"] = "Lookup"
    first["C2"] = 10
    first["B5"] = "Lookup"
    first["C5"] = 20
    second = workbook.create_sheet("Second")
    second["B2"] = "Lookup"
    second["C2"] = 30
    workbook.save(path)
    workbook.close()

    with ExcelReader.open(path) as reader:
        match = reader.get_value_from_anchor(
            "lookup",
            sheet="First",
            within="B1:C3",
        )

    assert match.value == 10
    assert match.anchor_cell.sheet == "First"


def test_missing_anchor_uses_stable_not_found_diagnostic(tmp_path: Path) -> None:
    path = tmp_path / "missing-anchor.xlsx"
    _save_anchor_workbook(path)

    with ExcelReader.open(path) as reader:
        matches = reader.find_values_from_anchor("missing", sheet="Form")

    assert not matches.matches
    assert matches.diagnostics[0].code is DiagnosticCode.ANCHOR_NOT_FOUND
    with pytest.raises(ExcelDataReaderError) as captured:
        matches.require_one()
    assert captured.value.diagnostics[0].code is DiagnosticCode.ANCHOR_NOT_FOUND


def test_anchor_lookup_preserves_formula_value_mode(tmp_path: Path) -> None:
    path = tmp_path / "formula-target.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "Lookup"
    sheet["B1"] = "=1+1"
    workbook.save(path)
    workbook.close()

    with ExcelReader.open(path, value_mode=ValueMode.BOTH) as reader:
        match = reader.get_value_from_anchor("lookup")

    assert match.value == FormulaValue("=1+1", None)
    assert match.value_cell.formula == "=1+1"


def test_offset_into_non_anchor_merged_cell_remains_physically_blank(tmp_path: Path) -> None:
    path = tmp_path / "merged-anchor.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.merge_cells("B2:C2")
    sheet["B2"] = "Lookup"
    workbook.save(path)
    workbook.close()

    with ExcelReader.open(path) as reader:
        match = reader.get_value_from_anchor("lookup")

    assert match.anchor_cell.address == "B2"
    assert match.value_cell.address == "C2"
    assert match.value is None


def test_text_anchor_lookup_does_not_match_non_text_cells(tmp_path: Path) -> None:
    path = tmp_path / "typed-anchors.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = 42
    sheet["A2"] = True
    sheet["A3"] = date(2026, 1, 2)
    sheet["A4"] = "#N/A"
    sheet["A5"] = "=1+1"
    sheet["A6"] = "42"
    sheet["B6"] = "text target"
    workbook.save(path)
    workbook.close()

    with ExcelReader.open(path) as reader:
        text_match = reader.get_value_from_anchor("42")
        non_text_results = tuple(
            reader.find_values_from_anchor(anchor)
            for anchor in ("true", "2026-01-02 00:00:00", "#N/A", "=1+1")
        )

    assert text_match.anchor_cell.address == "A6"
    assert text_match.value == "text target"
    assert all(not result.matches for result in non_text_results)
    assert all(
        result.diagnostics[0].code is DiagnosticCode.ANCHOR_NOT_FOUND for result in non_text_results
    )


@pytest.mark.parametrize("column_offset", [1, 16_383])
def test_blank_target_lookup_does_not_expand_apparent_sheet_bounds(
    tmp_path: Path,
    column_offset: int,
) -> None:
    path = tmp_path / f"blank-target-{column_offset}.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "Lookup"
    workbook.save(path)
    workbook.close()

    with ExcelReader.open(path, max_scan_cells=1) as reader:
        before = reader.inventory().sheets[0].dimension
        match = reader.get_value_from_anchor("lookup", column_offset=column_offset)
        after = reader.inventory().sheets[0].dimension
        repeated = reader.get_value_from_anchor("lookup", column_offset=column_offset)

    assert before == after == "A1:A1"
    assert match.value_cell.coordinate.column == 1 + column_offset
    assert match.value is None
    assert repeated.value is None
