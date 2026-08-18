from __future__ import annotations

import json
from pathlib import Path

import pytest
from openpyxl import Workbook

from excel_data_reader import (
    ANALYSIS_SCHEMA_VERSION,
    AnalysisOperation,
    AnalysisRequest,
    AnalysisStatus,
    DiagnosticCode,
    TableQuery,
    analyze_workbook,
)

ROOT = Path(__file__).parents[1]
WORKBOOKS = ROOT / "examples" / "workbooks"


def test_inventory_analysis_has_a_versioned_path_safe_contract() -> None:
    path = WORKBOOKS / "native_table.xlsx"

    response = analyze_workbook(path, AnalysisRequest.inventory(request_id="req-101"))
    payload = json.loads(response.to_json())

    assert response.status is AnalysisStatus.SUCCESS
    assert response.schema_version == ANALYSIS_SCHEMA_VERSION
    assert response.source_name == "native_table.xlsx"
    assert response.inventory is not None
    assert response.inventory.native_tables[0].name == "OrdersTable"
    assert payload["request_id"] == "req-101"
    assert str(path.parent) not in response.to_json()


def test_find_analysis_can_include_bounded_rows_and_typed_values() -> None:
    query = TableQuery(
        ("customer id", "amount"),
        optional_headers=("invoice date",),
        sheet="Scattered Orders",
    )
    request = AnalysisRequest.find_tables(query, include_rows=True, max_output_rows=2)

    response = analyze_workbook(WORKBOOKS / "scattered_headers.xlsx", request)
    payload = json.loads(response.to_json())

    assert response.status is AnalysisStatus.SUCCESS
    assert response.discovery is not None
    assert response.tables[0].total_row_count == 4
    assert response.tables[0].truncated is True
    assert len(response.tables[0].rows) == 2
    assert payload["tables"][0]["rows"][0]["cells"][2]["value"] == {
        "$type": "datetime",
        "value": "2026-03-01T00:00:00",
    }


def test_find_analysis_returns_no_match_without_throwing() -> None:
    request = AnalysisRequest.find_tables(TableQuery(("missing",), sheet="Scattered Orders"))

    response = analyze_workbook(WORKBOOKS / "scattered_headers.xlsx", request)

    assert response.status is AnalysisStatus.NO_MATCH
    assert response.discovery is not None
    assert not response.tables
    assert response.diagnostics[-1].code is DiagnosticCode.TABLE_NOT_FOUND


def test_find_analysis_reports_ambiguity_without_extracting_rows(tmp_path: Path) -> None:
    path = tmp_path / "ambiguous.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["ID", "Amount"])
    sheet.append(["A", 1])
    sheet["A5"] = "ID"
    sheet["B5"] = "Amount"
    sheet["A6"] = "B"
    sheet["B6"] = 2
    workbook.save(path)
    workbook.close()

    request = AnalysisRequest.find_tables(
        TableQuery(("id", "amount")),
        include_rows=True,
    )
    response = analyze_workbook(path, request)

    assert response.status is AnalysisStatus.AMBIGUOUS
    assert response.discovery is not None
    assert len(response.discovery.selected_matches) == 2
    assert not response.tables


def test_reader_diagnostics_are_mapped_to_error_responses() -> None:
    request = AnalysisRequest.find_tables(TableQuery(("id",), sheet="Absent"))

    response = analyze_workbook(WORKBOOKS / "native_table.xlsx", request)

    assert response.status is AnalysisStatus.ERROR
    assert response.diagnostics[0].code is DiagnosticCode.SHEET_NOT_FOUND


def test_analysis_request_rejects_inconsistent_operations() -> None:
    with pytest.raises(ValueError, match="requires a TableQuery"):
        AnalysisRequest(AnalysisOperation.FIND_TABLES)
