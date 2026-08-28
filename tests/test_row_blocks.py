from __future__ import annotations

from pathlib import Path
from typing import Any

import pytest
import xlwt
from openpyxl import Workbook
from openpyxl.styles import PatternFill

from excel_data_reader import (
    BlankRowSeparator,
    DiagnosticCode,
    ExcelDataReaderError,
    ExcelReader,
    RowBlockQuery,
    Severity,
    WorkbookFormat,
)


def _save_row_blocks(path: Path) -> None:
    """Create a workbook with shared headers and three blank-separated blocks.

    Args:
        path: Destination path for the temporary workbook.
    """

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Blocks"
    sheet["B6"] = "Value A"
    sheet["C6"] = "Value B"
    sheet["D6"] = "Value C"

    for row, label in ((8, "month1"), (9, "month2"), (11, "month3")):
        sheet.cell(row, 1, label)
        sheet.cell(row, 2, row * 10)
        sheet.cell(row, 3, row * 100)
        sheet.cell(row, 4, row * 1000)

    sheet["A12"].fill = PatternFill(fill_type="solid", fgColor="FFFF00")
    sheet["E13"] = "outside the scan range"

    for row, label in ((15, "year1"), (16, "year2")):
        sheet.cell(row, 1, label)
        sheet.cell(row, 2, row * 10)
        sheet.cell(row, 3, row * 100)
        sheet.cell(row, 4, row * 1000)

    sheet["A19"] = "metric1"
    sheet["B19"] = 190
    sheet["C19"] = 1900
    sheet["D19"] = 19000
    workbook.save(path)
    workbook.close()


def _save_variable_gap_blocks(path: Path) -> None:
    """Create four blocks separated by blank runs of varying lengths.

    Args:
        path: Destination path for the temporary workbook.
    """

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Blocks"
    sheet["B6"] = "Value A"
    sheet["C6"] = "Value B"
    sheet["D6"] = "Value C"
    for row, label in (
        (8, "month1"),
        (9, "month2"),
        (12, "year1"),
        (13, "year2"),
        (15, "metric1"),
        (17, "other1"),
    ):
        sheet.cell(row, 1, label)
        sheet.cell(row, 2, row * 10)
        sheet.cell(row, 3, row * 100)
        sheet.cell(row, 4, row * 1000)
    workbook.save(path)
    workbook.close()


def _query(*, min_rows: int = 2, scan_range: str = "A8:D25") -> RowBlockQuery:
    """Return the standard explicit row-block query used by tests.

    Args:
        min_rows: Consecutive blank rows required to split blocks.
        scan_range: Finite data rectangle inspected for populated rows.
    """

    return RowBlockQuery(
        sheet="Blocks",
        header_range="B6:D6",
        scan_range=scan_range,
        separator=BlankRowSeparator(min_rows=min_rows),
    )


def test_row_blocks_share_headers_and_preserve_short_internal_blank_rows(tmp_path: Path) -> None:
    path = tmp_path / "blocks.xlsx"
    _save_row_blocks(path)

    with ExcelReader.open(path) as reader:
        matches = reader.find_row_blocks(_query())
        first, second, third = matches.require_count(3)
        data = reader.extract_row_block(first)

    assert not matches.diagnostics
    assert [match.ordinal for match in matches.matches] == [1, 2, 3]
    assert [match.range for match in matches.matches] == ["A8:D11", "A15:D16", "A19:D19"]
    assert first.header_range.a1 == "B6:D6"
    assert first.separator_after is not None
    assert first.separator_after.a1 == "A12:D14"
    assert second.separator_after is not None
    assert second.separator_after.a1 == "A17:D18"
    assert third.separator_after is None
    assert [column.name for column in first.columns] == [
        "column_1",
        "Value A",
        "Value B",
        "Value C",
    ]
    assert first.columns[0].raw_header is None
    assert first.columns[0].header_coordinate is None
    assert first.columns[1].header_coordinate is not None
    assert first.columns[1].header_coordinate.a1 == "B6"
    assert len(data.rows) == 4
    assert data.rows[2].source_row == 10
    assert data.rows[2].values == (None, None, None, None)
    assert data.rows[0].cells[0].address == "A8"
    assert dict(data.records()[0]) == {
        "column_1": "month1",
        "Value A": 80,
        "Value B": 800,
        "Value C": 8000,
    }


def test_one_blank_row_can_be_declared_as_a_separator(tmp_path: Path) -> None:
    path = tmp_path / "one-blank.xlsx"
    _save_row_blocks(path)

    with ExcelReader.open(path) as reader:
        matches = reader.find_row_blocks(_query(min_rows=1))

    assert [match.range for match in matches.matches] == [
        "A8:D9",
        "A11:D11",
        "A15:D16",
        "A19:D19",
    ]
    assert matches.matches[0].separator_after is not None
    assert matches.matches[0].separator_after.a1 == "A10:D10"


def test_max_blocks_returns_requested_prefix_across_variable_separator_runs(
    tmp_path: Path,
) -> None:
    path = tmp_path / "variable-gaps.xlsx"
    _save_variable_gap_blocks(path)
    query = RowBlockQuery(
        sheet="Blocks",
        header_range="B6:D6",
        scan_range="A8:D25",
        separator=BlankRowSeparator(min_rows=1),
        max_blocks=3,
    )

    with ExcelReader.open(path) as reader:
        limited = reader.find_row_blocks(query)
        requested = limited.require_count(3)
        all_matches = reader.find_row_blocks(
            RowBlockQuery(
                sheet="Blocks",
                header_range="B6:D6",
                scan_range="A8:D25",
                separator=BlankRowSeparator(min_rows=1),
            )
        )

    assert [match.range for match in requested] == ["A8:D9", "A12:D13", "A15:D15"]
    assert requested[0].separator_after is not None
    assert requested[0].separator_after.a1 == "A10:D11"
    assert requested[1].separator_after is not None
    assert requested[1].separator_after.a1 == "A14:D14"
    assert requested[2].separator_after is None
    assert len(all_matches.matches) == 4
    assert not limited.diagnostics


def test_max_blocks_returns_available_prefix_when_request_cannot_be_filled(
    tmp_path: Path,
) -> None:
    path = tmp_path / "variable-gaps.xlsx"
    _save_variable_gap_blocks(path)
    query = RowBlockQuery(
        sheet="Blocks",
        header_range="B6:D6",
        scan_range="A8:D25",
        separator=BlankRowSeparator(min_rows=1),
        max_blocks=5,
    )

    with ExcelReader.open(path) as reader:
        matches = reader.find_row_blocks(query)

    assert len(matches.matches) == 4
    with pytest.raises(ExcelDataReaderError) as captured:
        matches.require_count(5)
    assert captured.value.diagnostics[0].code is DiagnosticCode.ROW_BLOCK_COUNT_MISMATCH


def test_column_names_apply_explicit_logical_names_and_preserve_source_metadata(
    tmp_path: Path,
) -> None:
    path = tmp_path / "named-columns.xlsx"
    _save_row_blocks(path)
    supplied_names = {"a": "row_label", "B": "amount"}
    query = RowBlockQuery(
        sheet="Blocks",
        header_range="B6:D6",
        scan_range="A8:D25",
        column_names=supplied_names,
        max_blocks=1,
    )
    supplied_names["a"] = "changed"

    with ExcelReader.open(path) as reader:
        match = reader.find_row_blocks(query).require_count(1)[0]
        data = reader.extract_row_block(match)

    assert dict(query.column_names) == {"A": "row_label", "B": "amount"}
    assert [column.name for column in data.columns] == [
        "row_label",
        "amount",
        "Value B",
        "Value C",
    ]
    assert data.columns[0].source_column == 1
    assert data.columns[0].raw_header is None
    assert data.columns[0].requested_header == "row_label"
    assert data.columns[0].header_coordinate is None
    assert data.columns[1].raw_header == "Value A"
    assert data.columns[1].requested_header == "amount"
    assert data.columns[1].header_coordinate is not None
    assert data.columns[1].header_coordinate.a1 == "B6"
    assert dict(data.records()[0]) == {
        "row_label": "month1",
        "amount": 80,
        "Value B": 800,
        "Value C": 8000,
    }


def test_column_names_must_identify_columns_inside_the_scan_range(tmp_path: Path) -> None:
    path = tmp_path / "outside-column-name.xlsx"
    _save_row_blocks(path)
    query = RowBlockQuery(
        sheet="Blocks",
        header_range="B6:D6",
        scan_range="A8:D25",
        column_names={"E": "outside"},
    )

    with ExcelReader.open(path) as reader, pytest.raises(ExcelDataReaderError) as captured:
        reader.find_row_blocks(query)

    diagnostic = captured.value.diagnostics[0]
    assert diagnostic.code is DiagnosticCode.INVALID_ROW_BLOCK_QUERY
    assert diagnostic.sheet == "Blocks"
    assert diagnostic.address == "A8:D25"


def test_populated_scan_bottom_warns_that_the_final_block_may_continue(tmp_path: Path) -> None:
    path = tmp_path / "scan-edge.xlsx"
    _save_row_blocks(path)

    with ExcelReader.open(path) as reader:
        matches = reader.find_row_blocks(_query(scan_range="A8:D19"))

    warning = matches.diagnostics[-1]
    assert warning.code is DiagnosticCode.ROW_BLOCK_SCAN_BOUNDARY_REACHED
    assert warning.severity is Severity.WARNING
    assert matches.matches[-1].diagnostics == (warning,)


def test_empty_scan_and_count_mismatch_have_stable_diagnostics(tmp_path: Path) -> None:
    path = tmp_path / "empty-scan.xlsx"
    _save_row_blocks(path)
    query = RowBlockQuery(
        sheet="Blocks",
        header_range="B6:D6",
        scan_range="A30:D40",
    )

    with ExcelReader.open(path) as reader:
        matches = reader.find_row_blocks(query)

    assert not matches.matches
    assert matches.diagnostics[0].code is DiagnosticCode.ROW_BLOCK_NOT_FOUND
    with pytest.raises(ExcelDataReaderError) as captured:
        matches.require_count(1)
    assert captured.value.diagnostics[0].code is DiagnosticCode.ROW_BLOCK_COUNT_MISMATCH
    assert captured.value.diagnostics[1].code is DiagnosticCode.ROW_BLOCK_NOT_FOUND


@pytest.mark.parametrize(
    ("header_range", "scan_range"),
    [
        ("B6:D7", "A8:D25"),
        ("B8:D8", "A8:D25"),
        ("B6:E6", "A8:D25"),
    ],
)
def test_row_block_query_rejects_inconsistent_ranges(
    tmp_path: Path,
    header_range: str,
    scan_range: str,
) -> None:
    path = tmp_path / "invalid-layout.xlsx"
    _save_row_blocks(path)
    query = RowBlockQuery(
        sheet="Blocks",
        header_range=header_range,
        scan_range=scan_range,
    )

    with ExcelReader.open(path) as reader, pytest.raises(ExcelDataReaderError) as captured:
        reader.find_row_blocks(query)

    assert captured.value.diagnostics[0].code is DiagnosticCode.INVALID_ROW_BLOCK_QUERY


def test_row_block_query_and_separator_validate_arguments() -> None:
    with pytest.raises(ValueError, match="min_rows"):
        BlankRowSeparator(min_rows=0)
    with pytest.raises(TypeError, match="min_rows"):
        BlankRowSeparator(min_rows=True)
    with pytest.raises(ValueError, match="sheet"):
        RowBlockQuery("", "B6:D6", "A8:D25")
    invalid_header: Any = 6
    with pytest.raises(TypeError, match="header_range"):
        RowBlockQuery("Blocks", invalid_header, "A8:D25")
    with pytest.raises(ValueError, match="max_blocks"):
        RowBlockQuery("Blocks", "B6:D6", "A8:D25", max_blocks=0)
    with pytest.raises(TypeError, match="max_blocks"):
        RowBlockQuery("Blocks", "B6:D6", "A8:D25", max_blocks=True)
    invalid_column_names: Any = []
    with pytest.raises(TypeError, match="column_names"):
        RowBlockQuery("Blocks", "B6:D6", "A8:D25", column_names=invalid_column_names)
    with pytest.raises(ValueError, match="column_names"):
        RowBlockQuery("Blocks", "B6:D6", "A8:D25", column_names={"A1": "row_label"})
    with pytest.raises(ValueError, match="column_names"):
        RowBlockQuery("Blocks", "B6:D6", "A8:D25", column_names={"A": " "})
    with pytest.raises(ValueError, match="duplicate"):
        RowBlockQuery(
            "Blocks",
            "B6:D6",
            "A8:D25",
            column_names={"A": "first", "a": "second"},
        )


def test_row_block_discovery_enforces_scan_and_candidate_limits(tmp_path: Path) -> None:
    path = tmp_path / "limits.xlsx"
    _save_row_blocks(path)

    with (
        ExcelReader.open(path, max_scan_cells=10) as reader,
        pytest.raises(ExcelDataReaderError) as captured,
    ):
        reader.find_row_blocks(_query())
    assert captured.value.diagnostics[0].code is DiagnosticCode.SCAN_LIMIT_EXCEEDED

    with ExcelReader.open(path, max_candidates=2) as reader:
        limited = reader.find_row_blocks(_query())
    assert len(limited.matches) == 2
    assert limited.diagnostics[0].code is DiagnosticCode.SCAN_LIMIT_EXCEEDED


def test_duplicate_shared_headers_remain_lossless_but_records_are_rejected(
    tmp_path: Path,
) -> None:
    path = tmp_path / "duplicate-headers.xlsx"
    _save_row_blocks(path)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Blocks"
    sheet["B6"] = "Amount"
    sheet["C6"] = "amount"
    sheet["A8"] = "month1"
    sheet["B8"] = 1
    sheet["C8"] = 2
    workbook.save(path)
    workbook.close()
    query = RowBlockQuery("Blocks", "B6:C6", "A8:C10")

    with ExcelReader.open(path) as reader:
        match = reader.find_row_blocks(query).require_count(1)[0]
        data = reader.extract_row_block(match)

    assert data.values == (("month1", 1, 2),)
    with pytest.raises(ExcelDataReaderError) as captured:
        data.records()
    assert captured.value.diagnostics[0].code is DiagnosticCode.DUPLICATE_HEADER


def test_row_blocks_use_the_legacy_xls_adapter(tmp_path: Path) -> None:
    path = tmp_path / "blocks.xls"
    workbook = xlwt.Workbook()
    sheet = workbook.add_sheet("Blocks")
    for column, value in enumerate(("Value A", "Value B", "Value C"), start=1):
        sheet.write(5, column, value)
    for row, label in ((7, "month1"), (8, "month2"), (11, "year1")):
        sheet.write(row, 0, label)
        sheet.write(row, 1, row * 10)
        sheet.write(row, 2, row * 100)
        sheet.write(row, 3, row * 1000)
    workbook.save(str(path))

    query = RowBlockQuery(
        "Blocks",
        "B6:D6",
        "A8:D15",
        column_names={"A": "row_label"},
    )
    with ExcelReader.open(path) as reader:
        matches = reader.find_row_blocks(query)
        first = reader.extract_row_block(matches.require_count(2)[0])
        workbook_format = reader.workbook_format

    assert workbook_format is WorkbookFormat.LEGACY_XLS
    assert [match.range for match in matches.matches] == ["A8:D9", "A12:D12"]
    assert first.columns[0].name == "row_label"
    assert first.values[0] == ("month1", 70.0, 700.0, 7000.0)
