from datetime import datetime
from pathlib import Path
from typing import Any, cast

import pytest
from openpyxl import Workbook

from excel_data_reader import (
    BodyPolicy,
    DiagnosticCode,
    ExcelDataReaderError,
    ExcelReader,
    FormulaValue,
    MultiLevelTableQuery,
    Rectangle,
    ValueMode,
    to_jsonable,
)
from tests.test_multilevel_discovery import repeated
from tests.test_multilevel_headers import layout


def workbook_path(tmp_path, kind="plain", suffix=".xlsx"):
    values, merges = repeated() if kind == "repeated" else layout(kind)[:2]
    if kind != "repeated":
        data_row = 5 if kind == "four" else 4
        values.update({(data_row, 1): 1001, (data_row, 2): 5, (data_row, 3): 7})
    path = tmp_path / ("headers" + suffix)
    if suffix == ".xls":
        xlwt = pytest.importorskip("xlwt")
        book = xlwt.Workbook()
        sheet = book.add_sheet("Report")
        for merge in merges:
            sheet.write_merge(
                merge.top - 1,
                merge.bottom - 1,
                merge.left - 1,
                merge.right - 1,
                values[merge.top, merge.left],
            )
        for (row, col), value in values.items():
            if not any(m.top <= row <= m.bottom and m.left <= col <= m.right for m in merges):
                sheet.write(row - 1, col - 1, value)
        book.save(str(path))
    else:
        book = Workbook()
        sheet = book.active
        sheet.title = "Report"
        for (row, col), value in values.items():
            sheet.cell(row, col, value)
        for merge in merges:
            sheet.merge_cells(merge.a1)
        book.save(path)
        book.close()
    return path


@pytest.mark.parametrize("suffix", [".xlsx", ".xls"])
@pytest.mark.parametrize("kind", ["plain", "merged", "wide", "mixed"])
def test_public_extraction_and_provenance(tmp_path, kind, suffix):
    path = workbook_path(tmp_path, kind, suffix)
    with ExcelReader.open(path) as reader:
        result = reader.find_multilevel_tables(MultiLevelTableQuery(("Sales Report",), 3))
        match = result.require_one()
        table = reader.extract_multilevel_table(match)
    assert table.rows[0].cells[0].address == "A4"
    assert table.records()[0]["Sales Report__col1"] == 1001
    source = table.columns[0].parts[0].source_coordinate
    assert source is not None and source.a1 == "A1"
    assert table.columns[0].parts[-1].source_bounds == Rectangle(2, 1, 3, 1)
    if kind in ("plain", "merged"):
        assert table.columns[-1].parts[-1].source_bounds == Rectangle(
            2, 3, 3 if kind == "merged" else 2, 3
        )
    if kind == "wide":
        assert table.columns[-1].name == "Sales Report__Store__col2"
        assert to_jsonable(table.columns[-1])["parts"][-1]["source_coordinate"] is None
    if suffix == ".xls":
        assert any(d.code is DiagnosticCode.LEGACY_XLS_LIMITED for d in result.diagnostics)


@pytest.mark.parametrize("suffix", [".xlsx", ".xls"])
def test_four_level_paths_through_public_reader(tmp_path, suffix):
    with ExcelReader.open(workbook_path(tmp_path, "four", suffix)) as reader:
        table = reader.extract_multilevel_table(
            reader.find_multilevel_tables(
                MultiLevelTableQuery((("Sales Report", "Region", "Store"),), 4)
            ).require_one()
        )
    assert [column.name for column in table.columns] == [
        "Sales Report__Region__Store__New",
        "Sales Report__Region__Store__Returning",
    ]
    assert table.rows[0].cells[0].address == "C5"
    assert table.values == ((7, None),)


def test_near_filter_runs_after_repeated_header_cap(tmp_path):
    with ExcelReader.open(workbook_path(tmp_path, "repeated")) as reader:
        result = reader.find_multilevel_tables(
            MultiLevelTableQuery(("Sales Report",), 3, near="A1")
        )
    assert result.require_one().data_end_row == 6


def test_typed_values_and_merge_placeholders(tmp_path):
    path = tmp_path / "typed.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.append(["Root"])
    sheet.merge_cells("A1:F1")
    sheet.append(["Date", "Bool", "Error", "Formula", "Merged", "Empty"])
    sheet.append([datetime(2026, 1, 2), True, "#DIV/0!", "=1+1", 25, None])
    sheet.merge_cells("E3:F3")
    sheet.row_dimensions[3].hidden = True
    sheet.column_dimensions["B"].hidden = True
    book.save(path)
    book.close()
    with ExcelReader.open(path, value_mode=ValueMode.BOTH) as reader:
        table = reader.extract_multilevel_table(
            reader.find_multilevel_tables(MultiLevelTableQuery(("Root",), 2)).require_one()
        )
    cells = table.rows[0].cells
    assert cells[0].is_date and cells[0].value == datetime(2026, 1, 2)
    assert cells[1].value is True and cells[1].hidden_column
    assert cells[2].value == "#DIV/0!" and cells[2].data_type == "e"
    assert cells[3].value == FormulaValue("=1+1", None)
    assert cells[4].value == 25 and cells[5].value is None
    assert all(cell.hidden_row for cell in cells)
    assert to_jsonable(table)["rows"][0]["cells"][0]["value"]["$type"] == "datetime"


def test_scope_limits_closed_reader_and_invalid_queries(tmp_path):
    path = workbook_path(tmp_path)
    with ExcelReader.open(path, max_scan_cells=2) as reader:
        result = reader.find_multilevel_tables(MultiLevelTableQuery(("Sales Report",), 3))
        assert not result.is_complete
        with pytest.raises(ExcelDataReaderError, match="SCAN_LIMIT_EXCEEDED"):
            result.require_one()
    with pytest.raises(ExcelDataReaderError, match="READER_CLOSED"):
        reader.find_multilevel_tables(MultiLevelTableQuery(("Sales Report",), 3))
    with ExcelReader.open(path) as reader:
        with pytest.raises(ExcelDataReaderError, match="INVALID_MULTILEVEL_QUERY"):
            reader.find_multilevel_tables(cast(Any, None))
        with pytest.raises(ExcelDataReaderError, match="SHEET_NOT_FOUND"):
            reader.find_multilevel_tables(
                MultiLevelTableQuery(("Sales Report",), 3, sheet="Missing")
            )
        result = reader.find_multilevel_tables(
            MultiLevelTableQuery(("Sales Report",), 3, within="A1:B4")
        )
        assert any(d.code is DiagnosticCode.MALFORMED_MULTILEVEL_HEADER for d in result.diagnostics)


def test_global_candidate_limit_preserves_only_completed_sheets(tmp_path):
    path = tmp_path / "sheets.xlsx"
    book = Workbook()
    for name in ("First", "Second"):
        sheet = book.active if name == "First" else book.create_sheet()
        sheet.title = name
        sheet.append(["Root"])
        sheet.append(["ID"])
        sheet.append([1])
    book.save(path)
    book.close()
    with ExcelReader.open(path, max_candidates=1) as reader:
        result = reader.find_multilevel_tables(MultiLevelTableQuery(("Root",), 2))
    assert not result.is_complete and [m.sheet for m in result.matches] == ["First"]


def test_legacy_checked_in_fixture_works_without_generator():
    path = Path(__file__).parents[1] / "examples/workbooks/legacy_scattered.xls"
    with ExcelReader.open(path) as reader:
        result = reader.find_multilevel_tables(MultiLevelTableQuery(("Customer ID", "Amount"), 1))
        table = reader.extract_multilevel_table(result.require_one())
    assert table.values[0] == ("L-001", 1250.0)


def test_cooperative_cancellation_during_extraction(tmp_path):
    cancel = False

    def checkpoint():
        if cancel:
            raise RuntimeError("cancelled")

    with ExcelReader.open(workbook_path(tmp_path), checkpoint=checkpoint) as reader:
        match = reader.find_multilevel_tables(
            MultiLevelTableQuery(("Sales Report",), 3)
        ).require_one()
        cancel = True
        with pytest.raises(RuntimeError, match="cancelled"):
            reader.extract_multilevel_table(match)


def test_explicit_scope_can_include_empty_header_rows_beyond_dimensions(tmp_path):
    path = tmp_path / "empty.xlsx"
    book = Workbook()
    sheet = book.active
    sheet["A1"] = "Root"
    sheet.merge_cells("A1:B1")
    book.save(path)
    book.close()
    with ExcelReader.open(path) as reader:
        result = reader.find_multilevel_tables(MultiLevelTableQuery(("Root",), 2, within="A1:B2"))
    match = result.require_one()
    assert match.is_empty and match.header_bounds.a1 == "A1:B2"
    assert [column.name for column in match.columns] == ["Root__col1", "Root__col2"]


def test_near_ties_keep_workbook_sheet_order_and_budget_counts_rejections(tmp_path):
    path = tmp_path / "two.xlsx"
    book = Workbook()
    for name in ("Z-first", "A-second"):
        sheet = book.active if name == "Z-first" else book.create_sheet()
        sheet.title = name
        sheet.append(["Root"])
        sheet.append(["ID"])
        sheet.append([1])
    book.save(path)
    book.close()
    with ExcelReader.open(path) as reader:
        result = reader.find_multilevel_tables(MultiLevelTableQuery(("Root",), 2, near="A1"))
        assert [match.sheet for match in result.matches] == ["Z-first", "A-second"]
        with pytest.raises(ExcelDataReaderError, match="AMBIGUOUS_TABLE"):
            result.require_one()
    with ExcelReader.open(path, max_candidates=1) as reader:
        result = reader.find_multilevel_tables(
            MultiLevelTableQuery(("Root",), 2, body=BodyPolicy.through_row(1))
        )
        assert not result.is_complete and result.candidate_count == 2
