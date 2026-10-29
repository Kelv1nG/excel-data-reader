"""Build the public multilevel-header workbook from entirely synthetic data."""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def _write_table(sheet, top: int, kind: str) -> None:
    """Write a three-row header and three data rows at the requested position."""

    width = 5 if kind == "mixed" else 4 if kind == "wide" else 3
    sheet.merge_cells(start_row=top, start_column=1, end_row=top, end_column=width)
    sheet.cell(top, 1, "Sales Report")
    sheet.merge_cells(start_row=top + 1, start_column=1, end_row=top + 2, end_column=1)
    sheet.cell(top + 1, 1, "col1")
    sheet.cell(top + 1, 2, "Online")
    sheet.cell(top + 2, 2, "New")
    sheet.cell(top + 1, 3, "Store")
    if kind == "merged":
        sheet.merge_cells(start_row=top + 1, start_column=3, end_row=top + 2, end_column=3)
    if kind in {"wide", "mixed"}:
        sheet.merge_cells(start_row=top + 1, start_column=3, end_row=top + 1, end_column=width)
    if kind == "mixed":
        sheet.cell(top + 2, 3, "New")
    for row in range(top + 3, top + 6):
        for column in range(1, width + 1):
            sheet.cell(row, column, row * 10 + column)
    for row in sheet.iter_rows(min_row=top, max_row=top + 2, min_col=1, max_col=width):
        for cell in row:
            cell.fill = PatternFill("solid", fgColor="17365D")
            cell.font = Font(bold=True, color="FFFFFF")
            cell.alignment = Alignment(horizontal="center", vertical="center")
    for column in range(1, width + 1):
        sheet.column_dimensions[get_column_letter(column)].width = 24
    sheet.sheet_view.showGridLines = False


def build_multilevel(output_dir: Path) -> Path:
    """Generate all discussion layouts as a standalone public example workbook."""

    book = Workbook()
    for name, kind in [
        ("Plain Store", "plain"),
        ("Merged Store", "merged"),
        ("Wide Store", "wide"),
        ("Mixed Leaves", "mixed"),
    ]:
        sheet = book.active if name == "Plain Store" else book.create_sheet()
        sheet.title = name
        _write_table(sheet, 1, kind)
    sheet = book.create_sheet("Four Levels")
    for span, value in [
        ("A1:D1", "Sales Report"),
        ("A2:D2", "Region"),
        ("A3:B3", "Online"),
        ("C3:D3", "Store"),
    ]:
        sheet.merge_cells(span)
        sheet[span.split(":")[0]] = value
    for col, value in enumerate(("New", "Returning", "New", "Returning"), 1):
        sheet.cell(4, col, value)
        sheet.cell(5, col, 50 + col)
    sheet = book.create_sheet("Repeated")
    _write_table(sheet, 1, "plain")
    _write_table(sheet, 7, "plain")
    output_dir.mkdir(parents=True, exist_ok=True)
    path = output_dir / "multilevel_tables.xlsx"
    book.save(path)
    book.close()
    return path


if __name__ == "__main__":
    print(build_multilevel(Path(__file__).parent / "workbooks"))
